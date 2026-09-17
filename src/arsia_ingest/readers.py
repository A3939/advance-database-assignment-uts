"""Read CSV and XLSX files into native rows.

CSV locators count logical records; XLSX locators use worksheet row numbers.
Both readers keep source order and duplicate rows. The value formats and blank
handling rules are described in docs/native-intake.md.

In ParseStats, ``records_seen`` counts observed data records or worksheet rows,
including skipped and invalid ones. ``raw_count`` counts yielded rows.
``blank_records_skipped`` counts CSV [] records; ``trailing_rows_skipped``
counts XLSX all-null rows at the end. ``header`` stores the observed header.
These counts are partial if reading stops early or fails. A parse is complete
only when the iterator finishes without an error.
"""
from __future__ import annotations

import csv
import json
import sys
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterator
from zipfile import BadZipFile
from xml.etree.ElementTree import ParseError

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from .models import IntakeError, NativeRow, ParseStats, ResourceSpec


def _error(spec: ResourceSpec, path: Path, code: str, message: str,
           locator: str | None = None, **details: object) -> IntakeError:
    return IntakeError(code, message, path=str(path), resource_id=spec.resource_id,
                       row_locator=locator, **details)


def _check_text(value: str, spec: ResourceSpec, path: Path,
                locator: str | None, column: str | None = None) -> None:
    if '\x00' in value:
        raise _error(spec, path, 'NUL_CHARACTER',
                     'NUL cannot be stored in PostgreSQL text or JSONB.', locator,
                     column=column)
    try:
        value.encode('utf-8', errors='strict')
    except UnicodeEncodeError as exc:
        raise _error(spec, path, 'INVALID_UTF8',
                     'The input must contain valid UTF-8 text.', locator,
                     column=column) from exc


def _validate_header(actual: list[object], spec: ResourceSpec, path: Path,
                     stats: ParseStats, locator: str) -> None:
    # Use repr for non-text cells so errors can be logged as JSON. Validate the
    # original types below, since converting them to text would hide bad headers.
    observed = [value if isinstance(value, str) else repr(value) for value in actual]
    stats.header = observed
    if any(not isinstance(value, str) for value in actual):
        raise _error(spec, path, 'HEADER_MISMATCH',
                     'Every header cell must be text and match the contract.',
                     locator, expected_header=list(spec.header), actual_header=observed)
    for value in actual:
        _check_text(value, spec, path, locator)  # type: ignore[arg-type]
    if len(actual) != len(set(actual)):
        raise _error(spec, path, 'DUPLICATE_HEADER',
                     'Duplicate headers would discard original fields.', locator,
                     actual_header=actual)
    if actual != list(spec.header):
        raise _error(spec, path, 'HEADER_MISMATCH',
                     'The complete ordered header must match the contract.', locator,
                     expected_header=list(spec.header), actual_header=actual)


def _raise_csv_limit() -> None:
    """Allow fields larger than 128 KiB without lowering an existing limit."""
    limit = sys.maxsize
    while True:
        try:
            if csv.field_size_limit() < limit:
                csv.field_size_limit(limit)
            return
        except OverflowError:
            limit //= 10


def _csv_rows(path: Path, spec: ResourceSpec, stats: ParseStats) -> Iterator[NativeRow]:
    _raise_csv_limit()
    logical = 0
    physical = 0
    header_pending = True
    try:
        # TextIOWrapper reads ahead. Validate UTF-8 one line at a time so an
        # invalid byte is reported against the right logical record.
        # surrogateescape holds those bytes until validated_lines rejects them.
        with path.open('r', encoding='utf-8-sig', errors='surrogateescape', newline='') as stream:
            def validated_lines() -> Iterator[str]:
                nonlocal physical
                for line in stream:
                    physical += 1
                    locator = 'csv:header' if header_pending else f'csv:{logical + 1}'
                    try:
                        line.encode('utf-8', errors='strict')
                    except UnicodeEncodeError as exc:
                        raise _error(spec, path, 'INVALID_UTF8',
                                     'The CSV must be valid UTF-8; only an initial BOM is stripped.',
                                     locator, physical_line=physical) from exc
                    yield line

            reader = csv.reader(validated_lines(), strict=True)
            try:
                actual = next(reader)
            except StopIteration:
                raise _error(spec, path, 'MISSING_HEADER', 'The input has no header.', 'csv:header')
            _validate_header(actual, spec, path, stats, 'csv:header')
            header_pending = False
            while True:
                try:
                    values = next(reader)
                except StopIteration:
                    break
                logical += 1
                stats.records_seen += 1
                locator = f'csv:{logical}'
                if values == []:
                    stats.blank_records_skipped += 1
                    continue
                if len(values) != len(spec.header):
                    raise _error(spec, path, 'ROW_WIDTH_MISMATCH',
                                 'Every CSV record must have exactly the declared number of fields.',
                                 locator, expected_columns=len(spec.header), actual_columns=len(values),
                                 physical_line=reader.line_num)
                for column, value in zip(spec.header, values):
                    _check_text(value, spec, path, locator, column)
                stats.raw_count += 1
                yield NativeRow(locator, dict(zip(spec.header, values)))
    except csv.Error as exc:
        if not header_pending:
            stats.records_seen += 1
        raise _error(spec, path, 'CSV_SYNTAX', str(exc),
                     'csv:header' if header_pending else f'csv:{logical + 1}',
                     physical_line=physical) from exc
    except IntakeError as exc:
        if exc.code == 'INVALID_UTF8' and not header_pending:
            stats.records_seen += 1
        raise


def _xlsx_locator(sheet: str, row: int) -> str:
    return json.dumps([sheet, row], ensure_ascii=False)


def _cell_value(cell: object, spec: ResourceSpec, path: Path,
                locator: str, column: str) -> str | None:
    value = cell.value  # type: ignore[attr-defined]
    kind = cell.data_type  # type: ignore[attr-defined]
    if kind in {'f', 'e', 'b'} or isinstance(value, bool):
        raise _error(spec, path, 'UNSUPPORTED_CELL_TYPE',
                     'Formula, error and Boolean cells require an explicit parser contract.',
                     locator, column=column, cell_type=kind)
    if value is None:
        return None
    if isinstance(value, str):
        _check_text(value, spec, path, locator, column)
        return value
    if isinstance(value, datetime):
        return value.isoformat(timespec='microseconds' if value.microsecond else 'seconds')
    if isinstance(value, time):
        return value.isoformat(timespec='microseconds' if value.microsecond else 'seconds')
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        try:
            number = Decimal(str(value))
        except InvalidOperation as exc:
            raise _error(spec, path, 'INVALID_NUMBER', 'The numeric cell is not a decimal value.',
                         locator, column=column) from exc
        if not number.is_finite():
            raise _error(spec, path, 'INVALID_NUMBER', 'Numeric cells must be finite.',
                         locator, column=column)
        return format(number, 'f')
    raise _error(spec, path, 'UNSUPPORTED_CELL_TYPE',
                 'The cell type is outside the fixed native parser contract.',
                 locator, column=column, cell_type=kind,
                 python_type=type(value).__name__)


def _blank_cell(cell: object) -> bool:
    # openpyxl can read an empty <v> as None while retaining its cell type.
    # Keep formula, Boolean and error cells for validation, even if they look empty.
    return cell.value is None and cell.data_type not in {'b', 'e', 'f'}  # type: ignore[attr-defined]


def _xlsx_rows(path: Path, spec: ResourceSpec, stats: ParseStats) -> Iterator[NativeRow]:
    sheet_name = spec.sheet
    if not sheet_name:
        raise _error(spec, path, 'INVALID_READER_SPEC', 'XLSX requires a specified worksheet.')
    # Archive filenames are hashes without extensions, so pass an open file.
    with path.open('rb') as binary_stream:
        workbook = load_workbook(binary_stream, read_only=True, data_only=False)
        try:
            if sheet_name not in workbook.sheetnames:
                raise _error(spec, path, 'SHEET_NOT_FOUND',
                             'The specified worksheet does not exist.', sheet=sheet_name)
            sheet = workbook[sheet_name]
            # Some workbooks declare A1:A1 despite having more rows in the XML.
            sheet.reset_dimensions()
            rows = sheet.iter_rows(min_row=1)
            try:
                try:
                    header_cells = next(rows)
                except StopIteration:
                    raise _error(spec, path, 'MISSING_HEADER', 'The worksheet has no header.',
                                 _xlsx_locator(sheet_name, 1))
                for index, cell in enumerate(header_cells, 1):
                    if cell.data_type in {'f', 'e', 'b'}:
                        raise _error(spec, path, 'HEADER_MISMATCH',
                                     'Every header cell must be literal text.',
                                     _xlsx_locator(sheet_name, 1), column_index=index,
                                     cell_type=cell.data_type)
                actual = [cell.value for cell in header_cells]
                while len(actual) > len(spec.header) and actual[-1] is None:
                    actual.pop()
                _validate_header(actual, spec, path, stats, _xlsx_locator(sheet_name, 1))
                width = len(spec.header)
                pending_count = 0
                pending_start = 0
                for row_number, cells in enumerate(rows, 2):
                    stats.records_seen += 1
                    locator = _xlsx_locator(sheet_name, row_number)
                    if all(_blank_cell(cell) for cell in cells):
                        if not pending_count:
                            pending_start = row_number
                        pending_count += 1
                        continue
                    # Later data means the pending blank rows must be kept.
                    for blank_number in range(pending_start, pending_start + pending_count):
                        stats.raw_count += 1
                        yield NativeRow(_xlsx_locator(sheet_name, blank_number),
                                        dict.fromkeys(spec.header))
                    pending_count = 0
                    extras = [(i + 1, cell) for i, cell in enumerate(cells[width:], width)
                              if not _blank_cell(cell)]
                    if extras:
                        raise _error(spec, path, 'ROW_WIDTH_MISMATCH',
                                     'A cell outside the declared header contains data.', locator,
                                     expected_columns=width, column_index=extras[0][0])
                    payload = {
                        name: _cell_value(cells[i], spec, path, locator, name) if i < len(cells) else None
                        for i, name in enumerate(spec.header)
                    }
                    stats.raw_count += 1
                    yield NativeRow(locator, payload)
                stats.trailing_rows_skipped += pending_count
            finally:
                rows.close()
        finally:
            workbook.close()


def iter_native_rows(path: Path, spec: ResourceSpec, stats: ParseStats) -> Iterator[NativeRow]:
    """Yield row locators and payloads from an archived input file.

    The caller adds file identity to each L1 record. Use fresh ``stats`` for each
    parse. If stopping early, close the iterator (or use ``contextlib.closing``)
    to release the file immediately.
    """
    if spec.header_row != 1 or not spec.header or any(
        not isinstance(name, str) or not name for name in spec.header
    ) or len(spec.header) != len(set(spec.header)):
        raise _error(spec, path, 'INVALID_READER_SPEC',
                     'The contract requires unique nonempty text headers on row 1.')
    if spec.format not in {'csv', 'xlsx'}:
        raise _error(spec, path, 'INVALID_READER_SPEC', 'Only csv and xlsx are supported.')
    if spec.format == 'csv' and spec.encoding not in {None, 'utf-8', 'utf-8-sig'}:
        raise _error(spec, path, 'INVALID_READER_SPEC', 'CSV encoding must be UTF-8.')
    try:
        if spec.format == 'csv':
            yield from _csv_rows(path, spec, stats)
        else:
            yield from _xlsx_rows(path, spec, stats)
    except (OSError, BadZipFile, InvalidFileException, ParseError, KeyError, ValueError) as exc:
        raise _error(spec, path, 'FILE_READ_ERROR',
                     f'The native file cannot be read: {exc}') from exc
