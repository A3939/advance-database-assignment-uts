"""Check the saved native observations before comparing official build results."""
from collections import Counter
import csv
import hashlib
import json
from pathlib import Path


EXPECTED_SHA256 = "71382aecb7969a023f48e52934bcb9873726a49b67807beea9c950184a01c609"
QA_RULES = ("QA01_INPUT", "QA02_RAW", "QA03_PROJECTED", "QA04_AUXILIARY",
            "QA05_SEMANTICS", "QA06_RECONCILIATION", "QA07_LOCATION")


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def native_year_counts(native_root, catalogue):
    """Read only native occurrence years, separately from the projection code."""
    from openpyxl import load_workbook

    by_year = Counter()
    resources = {row["resource_id"]: row for row in catalogue["resources"]}
    nsw = resources["official_nsw_crash"]
    workbook = load_workbook(native_root / Path(nsw["path"]).name, read_only=True, data_only=False)
    try:
        rows = workbook[nsw["sheet"]].iter_rows(values_only=True)
        headers = next(rows)
        index = headers.index("Year of crash")
        for row in rows:
            year = row[index]
            if year is not None and 2020 <= int(year) <= 2024:
                by_year[("official_nsw", int(year))] += 1
    finally:
        workbook.close()
    for resource, field, source in (("official_vic_accident", "ACCIDENT_DATE", "official_vic"),
                                    ("official_qld_crash", "Crash_Year", "official_qld")):
        entry = resources[resource]
        with (native_root / Path(entry["path"]).name).open(encoding=entry["encoding"], newline="") as stream:
            for row in csv.DictReader(stream):
                year = int(row[field][:4])
                if 2020 <= year <= 2024:
                    by_year[(source, year)] += 1
    return [{"source_id": source, "year": year, "crash_count": count}
            for (source, year), count in sorted(by_year.items())]


def load_expectations(project_root, native_root, manifest):
    """Use earlier native-file scans, never values measured by this build."""
    root, native_root = Path(project_root), Path(native_root)
    path = root / "config/official-expected-results-v1.json"
    assert digest(path) == EXPECTED_SHA256, "The reviewed expectations changed"
    expected = json.loads(path.read_text(encoding="utf-8"))
    assert expected["analysis"] == manifest["analysis"]
    assert expected["input_raw_count"] == 2118028
    evidence, references = {}, []
    for source in expected["sources"].values():
        for ref in source["evidence_refs"]:
            assert ref["path_base"] == "project_root"
            path = (root / ref["path"]).resolve()
            assert path.is_relative_to(root.resolve())
            assert digest(path) == ref["sha256"], ref["path"]
            evidence[ref["path"]] = json.loads(path.read_text(encoding="utf-8"))
            references.append(ref)
    nsw = evidence["docs/sources/evidence/official-decisions-2026-09-23/nsw-review.json"]["local_scan"]["analysis_totals"]
    vic = next(row for row in evidence["docs/sources/evidence/official-native-observations-2026-09-24.json"]["results"]
               if row["state"] == "VIC")["independent_native_totals"]
    qld = evidence["docs/sources/evidence/official-decisions-2026-09-23/qld-review.json"]["local_scan"]["analysis"]
    fields = ("crash_count", "fatal_crash_count", "fatality_count", "casualty_count", "canonical_unit_count")
    native_values = {
        "official_nsw": tuple(nsw[key] for key in ("crashes", "fatal_crashes", "fatalities", "casualties", "traffic_units")),
        "official_vic": tuple(vic[key] for key in ("crashes", "fatal_crashes", "fatalities", "casualties", "units")),
        "official_qld": (qld["rows"], qld["severity_counts"]["Fatal"], qld["fatality_sum"], qld["casualty_total_sum"], 0),
    }
    for source, values in native_values.items():
        assert tuple(expected["sources"][source][key] for key in fields) == values
    files = {row["resource_id"]: row for row in manifest["files"]}
    expected_files = {row["resource_id"]: row for source in expected["sources"].values() for row in source["input_files"]}
    catalogue = json.loads((root / "config/native-inputs.json").read_text(encoding="utf-8"))
    assert set(files) == set(expected_files) == {row["resource_id"] for row in catalogue["resources"]}
    native = []
    for entry in catalogue["resources"]:
        resource = entry["resource_id"]
        path = native_root / Path(entry["path"]).name
        sha256 = digest(path)
        assert sha256 == entry["expected_sha256"] == files[resource]["file_sha256"] == expected_files[resource]["file_sha256"]
        assert files[resource]["raw_count"] == expected_files[resource]["raw_count"]
        native.append({"resource_id": resource, "path": str(path), "sha256": sha256, "bytes": path.stat().st_size})
    year_counts = native_year_counts(native_root, catalogue)
    assert {(row["source_id"], row["year"]) for row in year_counts} == {
        (source, year) for source in expected["sources"] for year in range(2020, 2025)}
    for source, values in expected["sources"].items():
        assert sum(row["crash_count"] for row in year_counts if row["source_id"] == source) == values["crash_count"]
    expected["native_source_year_counts"] = year_counts
    return expected, {"native_files": native, "expectations_sha256": EXPECTED_SHA256,
                      "evidence_refs": references,
                      "native_source_year_counts": year_counts,
                      "basis": "Saved native-file metrics and a fresh occurrence-year scan, checked before running the platform; no new human-independent review claimed"}


def expected_qa_objects(manifest, expected):
    """The contract has 63 objects; all 16 location objects are limited here."""
    file_keys = {"file:" + ":".join((row["resource_id"], row["file_sha256"], row["parser_version"]))
                 for row in manifest["files"]}
    source_year = {f"source_year:{row['source_id']}:{row['year']}": row["crash_count"]
                   for row in expected["native_source_year_counts"]}
    groups = {
        "QA01_INPUT": file_keys, "QA02_RAW": file_keys,
        "QA03_PROJECTED": {"resource:" + name for name in ("official_nsw_crash", "official_nsw_traffic_unit",
                             "official_vic_accident", "official_vic_vehicle", "official_qld_crash")},
        "QA04_AUXILIARY": {"resource:" + name for name in ("official_nsw_traffic_unit", "official_vic_vehicle",
                             "official_vic_person", "official_vic_node")},
        "QA05_SEMANTICS": {"source:" + source for source in expected["sources"]},
        "QA06_RECONCILIATION": set(source_year), "QA07_LOCATION": set(source_year),
    }
    result = {(rule, key): ("pass", 0) for rule in QA_RULES for key in groups[rule] | {"batch"}}
    for key, count in {**source_year, "batch": sum(source_year.values())}.items():
        result[("QA07_LOCATION", key)] = ("limited", count)
    assert len(result) == 63
    return result
