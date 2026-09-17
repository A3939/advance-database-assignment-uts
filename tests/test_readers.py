"""Tests for native values, row positions and invalid CSV/XLSX input."""
from __future__ import annotations

import csv
import json
from datetime import date, datetime, time, timedelta
from pathlib import Path
from types import SimpleNamespace
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.styles import Font

from arsia_ingest.models import IntakeError, ParseStats, ResourceSpec
from arsia_ingest.readers import _cell_value, iter_native_rows


def spec(path: Path, header: tuple[str, ...] = ('id', 'value'), **kwargs: object) -> ResourceSpec:
    fmt = kwargs.pop('format', path.suffix[1:])
    return ResourceSpec('source', 'resource', 'crash', 'crash', path, fmt, header,
                        sheet='Data' if fmt == 'xlsx' else None, **kwargs)


def read(path: Path, contract: ResourceSpec | None = None):
    stats = ParseStats()
    return list(iter_native_rows(path, contract or spec(path), stats)), stats


def workbook(path: Path, rows: list[list[object]], sheet: str = 'Data') -> None:
    book = Workbook()
    page = book.active
    page.title = sheet
    for row in rows:
        page.append(row)
    book.save(path)
    book.close()


def rewrite_zip(path: Path, member: str, transform) -> None:
    with ZipFile(path) as archive:
        items = [(item.filename, archive.read(item.filename)) for item in archive.infolist()]
    with ZipFile(path, 'w', ZIP_DEFLATED) as archive:
        for name, data in items:
            archive.writestr(name, transform(data) if name == member else data)


def test_csv_bom_quotes_multiline_logical_locators_and_duplicate_rows(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_bytes(b'\xef\xbb\xbfid,value\r\n0001,"first\r\nsecond"\r\n\r\n0001,"first\r\nsecond"\r\n,""\r\n , Unknown \r\n')
    rows, stats = read(path)
    assert [row.row_locator for row in rows] == ['csv:1', 'csv:3', 'csv:4', 'csv:5']
    assert rows[0].payload == {'id': '0001', 'value': 'first\r\nsecond'}
    assert rows[1].payload == rows[0].payload
    assert rows[2].payload == {'id': '', 'value': ''}
    assert rows[3].payload == {'id': ' ', 'value': ' Unknown '}
    assert stats.raw_count == 4
    assert stats.records_seen == 5
    assert stats.blank_records_skipped == 1
    assert stats.trailing_rows_skipped == 0
    assert stats.header == ['id', 'value']


def test_csv_quoted_empty_and_whitespace_are_data_but_blank_records_are_skipped(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_bytes(b'id\r\n\r\n""\r\n \r\n0001\r\n\r\n')
    rows, stats = read(path, spec(path, ('id',)))
    assert [(row.row_locator, row.payload['id']) for row in rows] == [
        ('csv:2', ''), ('csv:3', ' '), ('csv:4', '0001')]
    assert (stats.raw_count, stats.records_seen, stats.blank_records_skipped) == (3, 5, 2)


def test_csv_cr_only_line_endings_and_embedded_bom_preserved(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_bytes('id,value\r0001,\ufeffNA\r0002,"a\rb"\r'.encode('utf-8'))
    rows, _ = read(path)
    assert rows[0].payload['value'] == '\ufeffNA'
    assert rows[1].payload['value'] == 'a\rb'
    assert rows[1].row_locator == 'csv:2'


def test_csv_does_not_filter_dates_or_interpret_tokens(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_text('id,value\n0001,1899\n0002,NA\n0003,-4\n0004,Unknown\n', encoding='utf-8')
    rows, _ = read(path)
    assert [row.payload['value'] for row in rows] == ['1899', 'NA', '-4', 'Unknown']


def test_csv_cell_larger_than_default_field_limit(tmp_path):
    path = tmp_path / 'input.csv'
    value = 'x' * 150_000
    path.write_text(f'id,value\n0001,{value}\n', encoding='utf-8')
    old_limit = csv.field_size_limit()
    csv.field_size_limit(131_072)
    try:
        rows, _ = read(path)
        assert rows[0].payload['value'] == value
    finally:
        csv.field_size_limit(old_limit)


@pytest.mark.parametrize(('body', 'code', 'locator'), [
    ('', 'MISSING_HEADER', 'csv:header'),
    ('id,id\n1,2\n', 'DUPLICATE_HEADER', 'csv:header'),
    ('value,id\n1,2\n', 'HEADER_MISMATCH', 'csv:header'),
    ('id\n1\n', 'HEADER_MISMATCH', 'csv:header'),
    ('id,value,extra\n1,2,3\n', 'HEADER_MISMATCH', 'csv:header'),
    ('id,value\n1\n', 'ROW_WIDTH_MISMATCH', 'csv:1'),
    ('id,value\n1,2,3\n', 'ROW_WIDTH_MISMATCH', 'csv:1'),
    ('id,value\n1,"unterminated\n', 'CSV_SYNTAX', 'csv:1'),
    ('id,value\n1,"ok"bad\n', 'CSV_SYNTAX', 'csv:1'),
    ('id,value\n1,abc\x00def\n', 'NUL_CHARACTER', 'csv:1'),
])
def test_csv_blocking_errors_include_file_and_record(tmp_path, body, code, locator):
    path = tmp_path / 'input.csv'
    path.write_text(body, encoding='utf-8')
    with pytest.raises(IntakeError) as caught:
        read(path)
    error = caught.value.as_dict()
    assert error['code'] == code
    assert error['row_locator'] == locator
    assert error['path'] == str(path)
    assert error['resource_id'] == 'resource'


def test_csv_invalid_utf8_is_localized_after_successful_record(tmp_path):
    path = tmp_path / 'input.csv'
    path.write_bytes(b'id,value\n0001,ok\n0002,\xff\n')
    stats = ParseStats()
    reader = iter_native_rows(path, spec(path), stats)
    assert next(reader).payload['id'] == '0001'
    with pytest.raises(IntakeError) as caught:
        next(reader)
    assert caught.value.code == 'INVALID_UTF8'
    assert caught.value.details['row_locator'] == 'csv:2'
    assert caught.value.details['physical_line'] == 3
    assert stats.raw_count == 1
    assert stats.records_seen == 2


def test_xlsx_values_iso_decimal_text_and_physical_positions(tmp_path):
    path = tmp_path / 'input.xlsx'
    header = ('id', 'integer', 'tiny', 'blank', 'date', 'datetime', 'time')
    book = Workbook()
    book.iso_dates = True
    sheet = book.active
    sheet.title = 'Data'
    sheet.append(header)
    sheet.append(['0001', 42, 1e-7, None, date(2020, 1, 15),
                  datetime(2020, 1, 15, 13, 4, 5, 123000), time(13, 4, 5)])
    sheet.append([None] * len(header))
    sheet.append(['0001', 42, 1e-7, None, date(2020, 1, 15),
                  datetime(2020, 1, 15, 13, 4, 5, 123000), time(13, 4, 5)])
    # Excel displays 0001 here, but the stored number is still 1.
    sheet.cell(5, 1, 1).number_format = '0000'
    sheet.cell(7, 1).font = Font(bold=True)
    sheet.cell(2, 12).font = Font(bold=True)
    book.save(path)
    book.close()
    rows, stats = read(path, spec(path, header))
    assert [json.loads(row.row_locator) for row in rows] == [
        ['Data', 2], ['Data', 3], ['Data', 4], ['Data', 5]]
    assert rows[0].payload == {
        'id': '0001', 'integer': '42', 'tiny': '0.0000001', 'blank': None,
        'date': '2020-01-15', 'datetime': '2020-01-15T13:04:05.123000', 'time': '13:04:05'}
    assert rows[1].payload == dict.fromkeys(header)
    assert rows[2].payload == rows[0].payload
    assert rows[3].payload['id'] == '1'
    assert (stats.raw_count, stats.records_seen, stats.trailing_rows_skipped) == (4, 6, 2)
    assert stats.blank_records_skipped == 0


def test_xlsx_default_excel_date_storage_yields_fixed_datetime(tmp_path):
    path = tmp_path / 'input.xlsx'
    workbook(path, [['id', 'value'], ['0001', date(2020, 1, 15)]])
    rows, _ = read(path)
    # openpyxl reads this Excel date as a datetime at midnight.
    assert rows[0].payload['value'] == '2020-01-15T00:00:00'


def test_xlsx_resets_incorrect_dimensions_and_does_not_skip_data(tmp_path):
    path = tmp_path / 'input.xlsx'
    workbook(path, [['id', 'value'], ['0001', 'x'], ['0002', 'y']])
    rewrite_zip(path, 'xl/worksheets/sheet1.xml',
                lambda data: data.replace(b'ref="A1:B3"', b'ref="A1:A1"'))
    rows, stats = read(path)
    assert len(rows) == 2
    assert stats.raw_count == 2
    assert rows[-1].payload == {'id': '0002', 'value': 'y'}


def test_xlsx_only_trailing_empty_rows_are_removed_and_extra_empty_header_cells_trimmed(tmp_path):
    path = tmp_path / 'input.xlsx'
    book = Workbook()
    sheet = book.active
    sheet.title = 'Data'
    sheet.append(['id', 'value'])
    sheet.cell(1, 8).font = Font(bold=True)
    sheet.cell(10, 1).font = Font(bold=True)
    book.save(path)
    book.close()
    rows, stats = read(path)
    assert rows == []
    assert stats.header == ['id', 'value']
    assert stats.records_seen == stats.trailing_rows_skipped == 9


@pytest.mark.parametrize(('value', 'cell_type'), [('=1+1', 'f'), (True, 'b'), ('#DIV/0!', 'e')])
def test_xlsx_unexpected_types_never_use_formula_cache_or_convert_bool(tmp_path, value, cell_type):
    path = tmp_path / 'input.xlsx'
    workbook(path, [['id', 'value'], ['0001', value]])
    with pytest.raises(IntakeError) as caught:
        read(path)
    assert caught.value.code == 'UNSUPPORTED_CELL_TYPE'
    assert caught.value.details['row_locator'] == '["Data", 2]'
    assert caught.value.details['column'] == 'value'
    assert caught.value.details['cell_type'] == cell_type


@pytest.mark.parametrize(('rows', 'code', 'row'), [
    ([['id', 'id']], 'DUPLICATE_HEADER', 1),
    ([['value', 'id']], 'HEADER_MISMATCH', 1),
    ([['id', 'value', 'extra']], 'HEADER_MISMATCH', 1),
    ([['id', 42]], 'HEADER_MISMATCH', 1),
    ([['id', date(2020, 1, 15)]], 'HEADER_MISMATCH', 1),
    ([['id', 'value'], ['0001', 'x', 'unexpected']], 'ROW_WIDTH_MISMATCH', 2),
])
def test_xlsx_schema_drift_blocks(tmp_path, rows, code, row):
    path = tmp_path / 'input.xlsx'
    workbook(path, rows)
    with pytest.raises(IntakeError) as caught:
        read(path)
    assert caught.value.code == code
    assert caught.value.details['row_locator'] == f'["Data", {row}]'
    json.dumps(caught.value.as_dict())


def test_xlsx_missing_sheet_blocks(tmp_path):
    path = tmp_path / 'input.xlsx'
    workbook(path, [['id', 'value']], sheet='Wrong')
    with pytest.raises(IntakeError, match='worksheet') as caught:
        read(path)
    assert caught.value.code == 'SHEET_NOT_FOUND'


@pytest.mark.parametrize('value', [float('nan'), float('inf'), float('-inf')])
def test_nonfinite_numeric_cells_block(tmp_path, value):
    path = tmp_path / 'input.xlsx'
    cell = SimpleNamespace(value=value, data_type='n')
    with pytest.raises(IntakeError) as caught:
        _cell_value(cell, spec(path), path, '["Data", 2]', 'value')
    assert caught.value.code == 'INVALID_NUMBER'


def test_duration_cell_requires_explicit_contract(tmp_path):
    path = tmp_path / 'input.xlsx'
    cell = SimpleNamespace(value=timedelta(hours=2), data_type='d')
    with pytest.raises(IntakeError) as caught:
        _cell_value(cell, spec(path), path, '["Data", 2]', 'value')
    assert caught.value.code == 'UNSUPPORTED_CELL_TYPE'


def test_reader_resources_close_when_generator_is_closed(tmp_path, monkeypatch):
    path = tmp_path / 'input.xlsx'
    workbook(path, [['id', 'value'], ['0001', 'x'], ['0002', 'y']])
    import arsia_ingest.readers as readers
    real_load = readers.load_workbook
    closed = []

    def tracked_load(*args, **kwargs):
        book = real_load(*args, **kwargs)
        real_close = book.close
        def close():
            closed.append(True)
            real_close()
        book.close = close
        return book

    monkeypatch.setattr(readers, 'load_workbook', tracked_load)
    iterator = iter_native_rows(path, spec(path), ParseStats())
    assert next(iterator).payload['id'] == '0001'
    iterator.close()
    assert closed == [True]


def test_invalid_archive_and_missing_file_are_intake_errors(tmp_path):
    path = tmp_path / 'bad.xlsx'
    path.write_bytes(b'not a zip archive')
    with pytest.raises(IntakeError) as caught:
        read(path)
    assert caught.value.code == 'FILE_READ_ERROR'
    with pytest.raises(IntakeError) as caught:
        read(tmp_path / 'missing.csv')
    assert caught.value.code == 'FILE_READ_ERROR'


def test_formula_header_cannot_pass_as_a_literal_matching_header(tmp_path):
    path = tmp_path / 'input.xlsx'
    workbook(path, [['id', '=value']])
    with pytest.raises(IntakeError) as caught:
        read(path, spec(path, ('id', '=value')))
    assert caught.value.code == 'HEADER_MISMATCH'


def test_structurally_invalid_zip_is_an_intake_error(tmp_path):
    path = tmp_path / 'input.xlsx'
    with ZipFile(path, 'w') as archive:
        archive.writestr('something.txt', 'not an Excel workbook')
    with pytest.raises(IntakeError) as caught:
        read(path)
    assert caught.value.code == 'FILE_READ_ERROR'


def test_xlsx_reads_extensionless_content_addressed_archive(tmp_path):
    original = tmp_path / 'input.xlsx'
    workbook(original, [['id', 'value'], ['0001', 'x']])
    archived = tmp_path / ('a' * 64)
    archived.write_bytes(original.read_bytes())
    rows, _ = read(archived, spec(archived, format='xlsx'))
    assert rows[0].payload == {'id': '0001', 'value': 'x'}


def test_csv_stream_closes_when_consumer_stops_early(tmp_path, monkeypatch):
    path = tmp_path / 'input.csv'
    path.write_text('id,value\n0001,x\n0002,y\n', encoding='utf-8')
    real_open = Path.open
    streams = []
    def tracked_open(self, *args, **kwargs):
        stream = real_open(self, *args, **kwargs)
        streams.append(stream)
        return stream
    monkeypatch.setattr(Path, 'open', tracked_open)
    iterator = iter_native_rows(path, spec(path), ParseStats())
    assert next(iterator).payload['id'] == '0001'
    assert not streams[0].closed
    iterator.close()
    assert streams[0].closed


@pytest.mark.parametrize('cell_type', ['b', 'e'])
@pytest.mark.parametrize(('column', 'expected_code'), [
    ('B', 'UNSUPPORTED_CELL_TYPE'), ('C', 'ROW_WIDTH_MISMATCH'),
])
def test_xlsx_empty_unsupported_cells_cannot_hide_as_blank_or_extra(tmp_path, cell_type, column, expected_code):
    path = tmp_path / 'input.xlsx'
    # Both cells have value=None after loading but retain their b/e data type.
    values = [None, True] if column == 'B' else [None, None, True]
    workbook(path, [['id', 'value'], values])
    original = f'<c r="{column}2" t="b"><v>1</v></c>'.encode()
    replacement = f'<c r="{column}2" t="{cell_type}"><v></v></c>'.encode()
    def empty_typed_value(data):
        assert original in data
        return data.replace(original, replacement)
    rewrite_zip(path, 'xl/worksheets/sheet1.xml', empty_typed_value)
    stats = ParseStats()
    with pytest.raises(IntakeError) as caught:
        list(iter_native_rows(path, spec(path), stats))
    assert caught.value.code == expected_code
    assert caught.value.details['row_locator'] == '["Data", 2]'
    assert stats.trailing_rows_skipped == 0
    assert stats.raw_count == 0
