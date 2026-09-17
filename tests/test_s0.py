"""Read S0 back from its native files and compare it with team contract 04."""
import csv
from decimal import Decimal
import importlib.util
import json
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook
import pytest

from arsia_ingest.pipeline import prepare


ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location("s0_generator", ROOT / "tools/create_s0_inputs.py")
s0 = importlib.util.module_from_spec(spec)
spec.loader.exec_module(s0)


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def read_native(directory):
    result = {}
    for resource in read_json(directory / "config.json")["resources"]:
        path = directory / resource["path"]
        if resource["format"] == "csv":
            with path.open(encoding="utf-8-sig", newline="") as stream:
                reader = csv.DictReader(stream)
                assert reader.fieldnames == resource["header"]
                result[resource["resource_id"]] = list(reader)
        else:
            book = load_workbook(path, read_only=True, data_only=False)
            try:
                values = book[resource["sheet"]].iter_rows(values_only=True)
                assert list(next(values)) == resource["header"]
                result[resource["resource_id"]] = [
                    {name: None if value is None else str(value)
                     for name, value in zip(resource["header"], row)} for row in values]
            finally:
                book.close()
    return result


def select(rows, *fields):
    return [tuple(row[field] for field in fields) for row in rows]


@pytest.fixture
def native(tmp_path):
    directory = tmp_path / "s0"
    s0.create_s0(directory)
    return directory, read_native(directory)


def test_s0_headers_rows_and_l1_match_native_files(native, tmp_path):
    directory, records = native
    expected_counts = {"syn_nsw_crash": 2, "syn_nsw_traffic_unit": 3,
                       "syn_vic_accident": 2, "syn_vic_vehicle": 3,
                       "syn_vic_person": 3, "syn_vic_node": 4, "syn_qld_crash": 2}
    assert {key: len(value) for key, value in records.items()} == expected_counts
    config = read_json(directory / "config.json")
    official_headers = read_json(ROOT / "config/native-inputs.json")["resources"]
    assert [r["header"] for r in config["resources"]] == [r["header"] for r in official_headers]
    assert sum(len(r["header"]) for r in config["resources"]) == 197
    result = prepare(directory / "config.json", tmp_path / "intake")
    assert (result["status"], result["raw_count"]) == ("prepared", 19)
    assert "batch_id" not in result and "input_fingerprint" not in result
    for item in result["files"]:
        path = Path(result["run_dir"]) / item["records_path"]
        rows = [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines()]
        assert [row["payload"] for row in rows] == records[item["resource_id"]]
        assert all(set(row) == {"source_id", "resource_id", "file_sha256", "parser_version",
                                "row_locator", "payload"} for row in rows)
        if item["resource_id"].startswith("syn_nsw"):
            assert [json.loads(row["row_locator"])[1] for row in rows] == list(range(2, len(rows) + 2))
        else:
            assert [row["row_locator"] for row in rows] == [f"csv:{i}" for i in range(1, len(rows) + 1)]


def test_s0_crash_values_follow_the_six_handwritten_examples(native):
    _, records = native
    assert select(records["syn_nsw_crash"], "Crash ID", "Year of crash", "Month of crash",
                  "Degree of crash - detailed", "No. killed", "No. seriously injured",
                  "No. moderately injured", "No. minor-other injured", "No. of traffic units involved") == [
        ("0001", "2020", "January", "F", "2", "1", "0", "0", "2"),
        ("0002", "2020", None, None, None, None, None, None, "1")]
    assert select(records["syn_vic_accident"], "ACCIDENT_NO", "ACCIDENT_DATE", "SEVERITY",
                  "NO_PERSONS_KILLED", "NO_PERSONS_INJ_2", "NO_PERSONS_INJ_3",
                  "NO_PERSONS_NOT_INJ", "NO_PERSONS", "NO_OF_VEHICLES", "NODE_ID") == [
        ("0001", "2020-01-15", "F", "1", "1", "0", "0", "2", "1", "NODE01"),
        ("0002", "2021-02-15", "I", "0", "1", "0", "0", "1", "2", "NODE02")]
    assert select(records["syn_qld_crash"], "Crash_Ref_Number", "Crash_Year", "Crash_Month",
                  "Crash_Severity", "Count_Casualty_Fatality", "Count_Casualty_Hospitalised",
                  "Count_Casualty_MedicallyTreated", "Count_Casualty_MinorInjury", "Count_Casualty_Total") == [
        ("0001", "2020", "January", "I", "0", "1", "0", "0", "1"),
        ("0002", "2021", "February", "N", "0", "0", "0", "0", "0")]
    assert all(row["Reporting year"] is None for row in records["syn_nsw_crash"])
    assert all(row["ACCIDENT_TIME"] == "" for row in records["syn_vic_accident"])


def test_s0_real_units_and_person_references_are_exact(native):
    _, records = native
    assert select(records["syn_nsw_traffic_unit"], "Crash ID", "Traffic unit ID", "TU type group") == [
        ("0001", "01", "CAR"), ("0001", "02", "CAR"), ("0002", "01", "CAR")]
    assert select(records["syn_vic_vehicle"], "ACCIDENT_NO", "VEHICLE_ID", "VEHICLE_TYPE") == [
        ("0001", "01", "CAR"), ("0002", "01", "CAR"), ("0002", "02", "CAR")]
    assert select(records["syn_vic_person"], "ACCIDENT_NO", "PERSON_ID", "VEHICLE_ID") == [
        ("0001", "P1", "01"), ("0001", "P2", "01"), ("0002", "P1", "")]
    assert not any("qld" in key and "unit" in key for key in records)
    assert all(int(value) >= 0 for row in records["syn_qld_crash"]
               for field, value in row.items() if field.startswith("Count_Unit_"))


def test_s0_keeps_equivalent_and_conflicting_node_observations(native):
    _, records = native
    nodes = records["syn_vic_node"]
    assert select(nodes, "ACCIDENT_NO", "NODE_ID") == [
        ("0001", "NODE01"), ("0001", "NODE01"), ("0002", "NODE02"), ("0002", "NODE02")]
    pairs = [(Decimal(r["LATITUDE"]), Decimal(r["LONGITUDE"])) for r in nodes]
    assert pairs == [(Decimal("-37.8"), Decimal("144.9")), (Decimal("-37.8"), Decimal("144.9")),
                     (Decimal("-37.9"), Decimal("145.0")), (Decimal("-38.0"), Decimal("145.1"))]
    assert nodes[0] != nodes[1]
    assert select(records["syn_nsw_crash"], "Latitude", "Longitude") == [
        ("-33.8600000", "151.2000000"), (None, None)]
    assert select(records["syn_qld_crash"], "Crash_Latitude", "Crash_Longitude") == [
        ("-27.4700000", "153.0200000"), ("-27.5000000", "153.0500000")]


def test_s0_contract_defines_the_synthetic_scope_and_unused_columns(native):
    directory, records = native
    contract = read_json(directory / "contract.json")
    assert contract["dataset_kind"] == "synthetic"
    assert contract["release_scope"] == "s0"
    assert contract["definition_version"] == "syn-1"
    assert contract["analysis"] == {"year_from": 2020, "year_to": 2024}
    assert contract["coverage"]["months"] == list(range(1, 13))
    assert [r["alias"] for r in contract["resources"]] == ["S1", "S2", "S3", "S4", "S5", "S6", "S7"]
    assert [(c["severity_code"], c["is_fatal_crash"]) for c in contract["severity"]["categories"]] == [
        ("F", True), ("I", False), ("N", False), ("__MISSING__", None)]
    assert contract["variant"]["test_status"] == "NOT_RUN"
    for resource in contract["resources"]:
        assert set(resource["used_fields"]).isdisjoint(resource["unused_fields"])
        assert set(resource["used_fields"] + resource["unused_fields"]) == set(resource["native"]["header"])
        for row in records[resource["resource_id"]]:
            assert all(row[field] in (None, "") for field in resource["unused_fields"])
    text = json.dumps(contract)
    assert "fatality_count_eligible" not in text and "casualty_count_eligible" not in text


def test_s0_regeneration_matches_the_shared_files_and_fixed_zip_times(tmp_path):
    directory = tmp_path / "copy"
    s0.create_s0(directory)
    stored = ROOT / "tests/fixtures/s0"
    assert {p.name: p.read_bytes() for p in directory.iterdir()} == {
        p.name: p.read_bytes() for p in stored.iterdir()}
    for path in directory.glob("*.xlsx"):
        with ZipFile(path) as archive:
            assert all(item.date_time == (2000, 1, 1, 0, 0, 0) for item in archive.infolist())
            assert b"2000-01-01T00:00:00Z" in archive.read("docProps/core.xml")


def test_s0_generator_refuses_to_replace_existing_work(tmp_path):
    directory = tmp_path / "existing"
    directory.mkdir()
    sentinel = directory / "keep.txt"
    sentinel.write_text("Keep this work.")
    with pytest.raises(ValueError, match="new or empty"):
        s0.create_s0(directory)
    assert list(directory.iterdir()) == [sentinel]
    assert sentinel.read_text() == "Keep this work."


@pytest.mark.parametrize("variant,code", [
    ("missing_file", "IO_ERROR"), ("bad_header", "HEADER_MISMATCH"), ("bad_hash", "FILE_HASH_MISMATCH")])
def test_s0_bad_native_inputs_fail_before_database_work(tmp_path, variant, code):
    config = s0.create_s0(tmp_path / "input", variant)
    result = prepare(config, tmp_path / "intake")
    assert result["status"] == "failed"
    assert result["errors"][0]["code"] == code
    assert not (Path(result["run_dir"]) / "records").exists()


@pytest.mark.parametrize("variant,rid,index,field,value", [
    ("duplicate_crash", "syn_nsw_crash", 2, "Crash ID", "0001"),
    ("orphan_unit", "syn_vic_vehicle", 2, "ACCIDENT_NO", "9999"),
    ("person_vehicle_99", "syn_vic_person", 2, "VEHICLE_ID", "99"),
    ("invalid_date", "syn_vic_accident", 1, "ACCIDENT_DATE", "2021-02-30"),
    ("negative_count", "syn_vic_accident", 0, "NO_PERSONS_KILLED", "-1"),
    ("undefined_category", "syn_vic_accident", 1, "SEVERITY", "X"),
    ("invalid_coordinate", "syn_qld_crash", 0, "Crash_Latitude", "-91"),
    ("revised_n1", "syn_nsw_crash", 0, "No. killed", "3"),
])
def test_semantic_variants_reach_sql_with_their_original_values(tmp_path, variant, rid, index, field, value):
    directory = tmp_path / "input"
    config = s0.create_s0(directory, variant)
    assert read_native(directory)[rid][index][field] == value
    result = prepare(config, tmp_path / "intake")
    assert result["status"] == "prepared"
    assert result["raw_count"] == (20 if variant == "duplicate_crash" else 19)
    assert read_json(directory / "contract.json")["variant"]["test_status"] == "NOT_RUN"


def test_deletion_variants_have_the_same_data_but_different_change_evidence(tmp_path):
    first, second = tmp_path / "explained", tmp_path / "unexplained"
    s0.create_s0(first, "delete_q2")
    s0.create_s0(second, "delete_q2_unexplained")
    assert read_native(first) == read_native(second)
    assert len(read_native(first)["syn_qld_crash"]) == 1
    assert "snapshot_change" in read_json(first / "contract.json")
    assert "snapshot_change" not in read_json(second / "contract.json")
    for directory in (first, second):
        result = prepare(directory / "config.json", tmp_path / f"intake-{directory.name}")
        assert (result["status"], result["raw_count"]) == ("prepared", 18)


def test_unknown_crs_changes_only_the_declared_location_basis(tmp_path):
    baseline, variant = tmp_path / "baseline", tmp_path / "unknown"
    s0.create_s0(baseline)
    s0.create_s0(variant, "unknown_crs")
    assert read_native(baseline) == read_native(variant)
    contracts = read_json(variant / "contract.json")["resources"]
    qld = next(r for r in contracts if r["resource_id"] == "syn_qld_crash")
    assert qld["mapping"]["location"]["crs"] is None
    result = prepare(variant / "config.json", tmp_path / "intake")
    assert result["status"] == "prepared"
