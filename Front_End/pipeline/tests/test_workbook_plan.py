import copy
from datetime import datetime
import hashlib
import json

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904
import pytest

from arsia_pipeline.errors import ImportCancelled, UnsupportedCapability, ValidationFailure
from arsia_pipeline.intakereaders import detect_tables, iter_table
from arsia_pipeline.intake_tools import IntakeTools
from arsia_pipeline.trusted_qa import validate_contract
from test_canonical_v2 import contract


def workbook(tmp_path, extra=False):
    book = Workbook()
    sheet = book.active
    sheet.title = 'Records'
    sheet.append(['Export catalogue: Event records'])
    sheet.append(['(Data current to September 2026)'])
    sheet.append(['Note: Unknown values are retained.'])
    sheet.append([])
    sheet.append(['ID', 'DATE', 'SEVERITY', 'DEATHS', 'INJURED'])
    sheet.append(['001', '30/04/2024 13:20', 'F', 2, 3])
    sheet.append([])
    sheet.append(['002', '01/05/2024 13:20', 'U', 0, 1])
    if extra:
        hidden = book.create_sheet('Other facts')
        hidden.sheet_state = 'hidden'
        hidden.append([])
        hidden.append(['ID', 'Total'])
        hidden.append(['A', 10])
        book.create_sheet('Empty')
    path = tmp_path / 'opaque'
    book.save(path)
    return book, file_for(path)


def file_for(path):
    return {'id': 'file1', 'name': path.name, 'path': str(path), 'size': path.stat().st_size,
            'sha256': hashlib.sha256(path.read_bytes()).hexdigest()}


def test_prefix_plan_and_physical_row_conservation(tmp_path):
    _, file = workbook(tmp_path)
    table = detect_tables(file)[0]
    assert table['header_row'] == 5
    assert table['parser_plan']['input_sha256'] == file['sha256']
    assert table['parser_plan']['data_end_row'] is None
    assert len(table['parser_plan']['prefix']) == 4
    assert [p['row'] for p in table['parser_plan']['prefix']] == [1, 2, 3, 4]
    assert 'Unknown values' not in json.dumps(table)  # No arbitrary cell prose in model inspection.
    rows = list(iter_table(file, table))
    assert [r[0] for r in rows] == ['["Records",6]', '["Records",7]', '["Records",8]']
    assert rows[1][1] == dict.fromkeys(table['header'])
    assert rows[-1][1]['ID'] == '002'


def test_profile_uses_the_same_plan_without_leaking_records(tmp_path):
    _, file = workbook(tmp_path)
    tools = IntakeTools([file], tmp_path / 'task')
    inspection = tools.inspect_bundle()
    table = inspection['resources'][0]['tables'][0]
    result = tools.profile_dataset('file1', 'Records', {'header_row': 5, 'parser_plan': table['parser_plan']})
    assert result['row_count'] == 3
    assert '30/04/2024 13:20' not in json.dumps(inspection)


@pytest.mark.parametrize('key,value', [('header_row', 6), ('header_row', True), ('header', ['not original']),
                                     ('format', 'csv'), ('skipRows', 6), ('end_row', 6)])
def test_hints_cannot_select_another_header_or_filter_records(tmp_path, key, value):
    _, file = workbook(tmp_path)
    spec = detect_tables(file)[0]
    spec[key] = value
    with pytest.raises((ValidationFailure, UnsupportedCapability)):
        list(iter_table(file, spec))


def test_stale_plan_and_changed_input_are_rejected(tmp_path):
    book, file = workbook(tmp_path)
    spec = detect_tables(file)[0]
    book.active['A8'] = 'changed'
    book.save(file['path'])
    with pytest.raises(ValidationFailure, match='input hash'):
        detect_tables(file)
    fresh = file_for(__import__('pathlib').Path(file['path']))
    with pytest.raises(ValidationFailure, match='observed layout'):
        list(iter_table(fresh, spec))
    changed = copy.deepcopy(detect_tables(fresh)[0])
    changed['parser_plan']['prefix'][0]['values_sha256'] = '0' * 64
    with pytest.raises(ValidationFailure, match='observed layout'):
        list(iter_table(fresh, changed))


def test_hidden_facts_and_empty_sheets_remain_in_inventory(tmp_path):
    _, file = workbook(tmp_path, extra=True)
    tables = detect_tables(file, {'sheet': 'Records', 'header_row': 5})
    assert [(t['sheet'], t['header_row'], t['sheet_state'], t['empty']) for t in tables] == [
        ('Records', 5, 'visible', False), ('Other facts', 2, 'hidden', False), ('Empty', None, 'visible', True)]
    assert list(iter_table(file, tables[1])) == [('["Other facts",3]', {'ID': 'A', 'Total': '10'})]
    c = contract()
    c['resources'][0]['table'] = tables[0]
    with pytest.raises(ValidationFailure, match='assigned a grain exactly once'):
        validate_contract(c, [file])
    with pytest.raises(ValidationFailure, match='empty worksheet'):
        list(iter_table(file, tables[2]))


@pytest.mark.parametrize('prefix', [
    [['Events', 'Count'], ['A', 2], []], # An earlier fact region cannot be skipped.
    [['Catalogue: Events'], ['not a labelled annotation'], []],
    [['Catalogue: Events'], ['(Annotations)']], # No separator.
    [['Catalogue: Events'], ['(Annotations)'], [10], []],
    [['Catalogue: Events'], ['(Annotations)'], ['Note: text', 'unclassified'], []],
])
def test_data_or_ambiguous_preamble_cannot_be_hidden(tmp_path, prefix):
    book = Workbook()
    for row in prefix + [['ID', 'Year'], ['A', 2024]]:
        book.active.append(row)
    path = tmp_path / 'ambiguous.xlsx'
    book.save(path)
    with pytest.raises((ValidationFailure, UnsupportedCapability)):
        # Pinning the later valid-looking string row never authorizes the skip.
        list(iter_table(path, {'format': 'xlsx', 'header_row': len(prefix) + 1}))


def test_late_nonempty_sheet_is_not_omitted(tmp_path):
    book = Workbook()
    book.active['A70'] = 'ID'
    book.active['A71'] = 'FACT'
    path = tmp_path / 'late.xlsx'
    book.save(path)
    with pytest.raises(UnsupportedCapability, match='beyond'):
        detect_tables(path)


def test_native_epoch_and_exact_header_names_are_preserved(tmp_path):
    book = Workbook()
    book.epoch = CALENDAR_MAC_1904
    book.active.append([' ID ', 'id', 'ＤＡＴＥ'])
    book.active.append(['A', 'B', datetime(2024, 2, 29)])
    path = tmp_path / 'calendar.xlsx'
    book.save(path)
    spec = detect_tables(path)[0]
    assert spec['header'] == [' ID ', 'id', 'ＤＡＴＥ']
    assert spec['parser_plan']['epoch'] == '1904-01-01T00:00:00'
    assert list(iter_table(path))[0][1]['ＤＡＴＥ'] == '2024-02-29T00:00:00'


@pytest.mark.parametrize('cell,value', [('A1', '=1+1'), ('A6', '=1+1'), ('A9', '#REF!')])
def test_formula_errors_in_prefix_data_or_tail_cannot_disappear(tmp_path, cell, value):
    book, file = workbook(tmp_path)
    book.active[cell] = value
    book.save(file['path'])
    with pytest.raises(ValidationFailure, match='Formula/error'):
        list(iter_table(file_for(__import__('pathlib').Path(file['path']))))


def test_cancel_interrupts_layout_and_input_hashing(tmp_path):
    _, file = workbook(tmp_path)
    def cancelled():
        raise ImportCancelled()
    with pytest.raises(ImportCancelled):
        detect_tables(file, check_cancelled=cancelled)
