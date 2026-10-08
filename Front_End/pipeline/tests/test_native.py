"""Adversarial native-profile boundary tests; tiny rows are labelled fixtures.

These unit tests do not pretend that synthetic fixtures have the pinned hashes
or supply official publication acceptance. Full-source API tests are separate.
"""
import csv
import hashlib
import json
from pathlib import Path
import sqlite3
import zipfile

import pytest
from openpyxl import Workbook

from arsia_pipeline.errors import ImportCancelled, NeedsInput, ValidationFailure
from arsia_pipeline.native import (_crashes, _database, _export, _profile, _references,
                                   _relationships, _same_cases, identify_bundle, process_native)
from arsia_pipeline.readers import inspect_file, iter_rows


def file_info(path):
    return {"id": path.name, "name": path.name, "path": str(path), "size": path.stat().st_size,
            "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def write_csv(path, header, rows=()):
    with path.open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(header)
        writer.writerows(rows)
    return file_info(path)


def native_header(role, state):
    return next(r["header"] for r in _references()[1]["native-inputs.json"]["resources"]
                if r["resource_role"] == role and r["source_id"] == f"official_{state.lower()}")


def qld_row(**overrides):
    row = dict.fromkeys(native_header("crash", "QLD"), "")
    row.update(Crash_Ref_Number="fixture-001", Crash_Year="2023", Crash_Month="January", Crash_Severity="Fatal",
               Count_Casualty_Fatality="2", Count_Casualty_Hospitalised="1", Count_Casualty_MedicallyTreated="0",
               Count_Casualty_MinorInjury="0", Count_Casualty_Total="3")
    row.update(overrides)
    return row


def selected_qld(tmp_path, row):
    file = write_csv(tmp_path / "renamed.csv", list(row), [list(row.values())])
    return {"file": file, "table": inspect_file(file)[0], "spec": {"resource_id": "official_qld_crash"}}


def test_frozen_reference_files_retain_upstream_hashes():
    upstream, refs = _references()
    assert upstream["commit"] == "39fd9032dd792ab5f350e1eaf66697653a263272"
    assert len(refs["native-inputs.json"]["resources"]) == 7


def test_identification_uses_headers_not_filename(tmp_path):
    file = write_csv(tmp_path / "definitely_not_queensland.csv", native_header("crash", "QLD"))
    bundle = identify_bundle([file])
    assert bundle["state"] == "QLD"
    assert set(bundle["files_by_role"]) == {"crash"}


def test_new_bytes_cannot_inherit_old_native_approval(tmp_path):
    file = write_csv(tmp_path / "qld.csv", native_header("crash", "QLD"))
    with pytest.raises(NeedsInput, match="new bytes"):
        process_native([file], tmp_path / "attempt", {}, lambda *a, **k: None, lambda: None)
    assert not (tmp_path / "attempt" / "canonical.jsonl").exists()


def test_changed_native_headers_do_not_fall_through_to_generic(tmp_path):
    file = write_csv(tmp_path / "qld.csv", native_header("crash", "QLD") + ["new_field"])
    with pytest.raises(NeedsInput, match="unreviewed"):
        identify_bundle([file])


def test_vic_requires_all_four_files(tmp_path):
    file = write_csv(tmp_path / "accidents.csv", native_header("crash", "VIC"))
    with pytest.raises(NeedsInput) as error:
        identify_bundle([file])
    assert error.value.details["missing_roles"] == ["node", "person", "vehicle"]


def test_duplicate_and_mixed_source_roles_block(tmp_path):
    qld = write_csv(tmp_path / "one.csv", native_header("crash", "QLD"))
    same = write_csv(tmp_path / "two.csv", native_header("crash", "QLD"))
    with pytest.raises(NeedsInput, match="same source role"):
        identify_bundle([qld, same])
    vic = write_csv(tmp_path / "vic.csv", native_header("crash", "VIC"))
    with pytest.raises(NeedsInput, match="one state's"):
        identify_bundle([qld, vic])


def test_notes_and_dictionary_are_evidence_not_data(tmp_path):
    qld = write_csv(tmp_path / "one.csv", native_header("crash", "QLD"))
    dictionary = write_csv(tmp_path / "fields.csv", ["field", "description"], [["Crash_Year", "Occurrence year"]])
    notes = tmp_path / "origin.txt"
    notes.write_text("Synthetic unit-test source note.")
    bundle = identify_bundle([qld, dictionary, file_info(notes)])
    assert len(bundle["supplements"]) == 2


def test_unknown_extra_data_is_not_silently_ignored(tmp_path):
    qld = write_csv(tmp_path / "one.csv", native_header("crash", "QLD"))
    other = write_csv(tmp_path / "other.csv", ["id", "year"], [["x", "2020"]])
    with pytest.raises(NeedsInput, match="Unrecognized data"):
        identify_bundle([qld, other])


def test_csv_locator_counts_logical_data_records(tmp_path):
    file = write_csv(tmp_path / "fixture.csv", ["id", "note"], [["a", "two\nlines"], ["b", "next"]])
    rows = list(iter_rows(file))
    assert [locator for locator, _ in rows] == ["csv:1", "csv:2"]
    assert rows[0][1]["note"] == "two\nlines"


def test_csv_bad_width_and_duplicate_header_block(tmp_path):
    bad = write_csv(tmp_path / "width.csv", ["id", "value"], [["one"]])
    with pytest.raises(ValidationFailure, match="width"):
        list(iter_rows(bad))
    duplicate = write_csv(tmp_path / "headers.csv", ["id", "id"])
    with pytest.raises(ValidationFailure, match="Duplicate"):
        inspect_file(duplicate)


def workbook(tmp_path, rows):
    path = tmp_path / "fixture.xlsx"
    book = Workbook()
    for row in rows:
        book.active.append(row)
    book.save(path)
    book.close()
    return path


def test_xlsx_formulas_are_not_executed_or_cached(tmp_path):
    path = workbook(tmp_path, [["id", "value"], ["x", "=1+1"]])
    with pytest.raises(ValidationFailure, match="formulas"):
        list(iter_rows(file_info(path)))


def test_xlsx_interior_blank_rows_are_preserved_tail_is_skipped(tmp_path):
    path = workbook(tmp_path, [["id", "value"], ["a", 1], [None, None], ["b", 2], [None, None]])
    rows = list(iter_rows(file_info(path)))
    assert [json.loads(locator)[1] for locator, _ in rows] == [2, 3, 4]
    assert rows[1][1] == {"id": None, "value": None}


def test_xlsx_ignores_incorrect_dimension_hint(tmp_path):
    path = workbook(tmp_path, [["id", "value"], ["a", 1], ["b", 2]])
    with zipfile.ZipFile(path) as original:
        members = {name: original.read(name) for name in original.namelist()}
    members["xl/worksheets/sheet1.xml"] = members["xl/worksheets/sheet1.xml"].replace(b'ref="A1:B3"', b'ref="A1:A1"')
    with zipfile.ZipFile(path, "w") as updated:
        for name, content in members.items():
            updated.writestr(name, content)
    assert len(list(iter_rows(file_info(path)))) == 2


def test_cancellation_is_propagated(tmp_path):
    file = write_csv(tmp_path / "x.csv", ["id"], [[str(i)] for i in range(1001)])
    def cancelled():
        raise ImportCancelled()
    with pytest.raises(ImportCancelled):
        list(iter_rows(file, check_cancelled=cancelled))


@pytest.mark.parametrize("overrides, message", [
    ({"Crash_Severity": "assumed serious"}, "severity"),
    ({"Count_Casualty_Total": "4"}, "total differs"),
    ({"Crash_Ref_Number": "  "}, "native key"),
    ({"Count_Casualty_Fatality": "-1"}, "nonnegative"),
    ({"Crash_Month": "month 13"}, "month"),
])
def test_qld_rejects_ambiguous_or_inconsistent_source_values(tmp_path, overrides, message):
    _, definitions, native, _ = _profile("QLD", _references()[1])
    db = _database(tmp_path / "index.sqlite")
    try:
        with pytest.raises(ValidationFailure, match=message):
            _crashes(db, selected_qld(tmp_path, qld_row(**overrides)), "QLD", native, definitions,
                     lambda *a, **k: None, lambda: None)
    finally:
        db.close()


def test_unknown_counts_stay_null_and_fatal_crash_is_not_fatalities(tmp_path):
    _, definitions, native, _ = _profile("QLD", _references()[1])
    selected = selected_qld(tmp_path, qld_row(Count_Casualty_Hospitalised=""))
    db = _database(tmp_path / "index.sqlite")
    try:
        _crashes(db, selected, "QLD", native, definitions, lambda *a, **k: None, lambda: None)
        bundle = {"source_id": "official_qld", "state": "QLD", "files_by_role": {"crash": selected}}
        result = _export(db, tmp_path, bundle, definitions, lambda: None)
        assert result[0]["crash_count"] == 1
        assert result[0]["fatal_crash_count"] == 1
        assert result[0]["fatalities"] == 2
        assert result[0]["casualties"] is None
        assert result[4]["status"] == "unavailable"
    finally:
        db.close()


def test_nsw_relationship_checks_all_years_not_only_analysis(tmp_path):
    db = _database(tmp_path / "index.sqlite")
    try:
        db.execute("INSERT INTO crash VALUES ('old',2019,1,NULL,'FATAL',1,1,1,2,NULL,NULL,'fixture')")
        db.execute("INSERT INTO unit VALUES ('old','1','Pedestrian','fixture')")
        with pytest.raises(ValidationFailure, match="declared traffic-unit"):
            _relationships(db, "NSW", _references()[1], {}, lambda: None)
    finally:
        db.close()


def test_vic_case_counts_alone_do_not_approve_changed_identity():
    with pytest.raises(ValidationFailure, match="exact reviewed case"):
        _same_cases([("new-key", 39)], [("pinned-key", 39)], "fixture")


def test_vic_retains_units_without_enabling_vehicle_reporting(tmp_path):
    db = _database(tmp_path / "index.sqlite")
    _, definitions, _, _ = _profile("VIC", _references()[1])
    try:
        db.execute("INSERT INTO crash VALUES ('fixture',2023,1,'2023-01-01','1',1,1,1,1,1,'node','csv:1')")
        db.execute("INSERT INTO unit VALUES ('fixture','A','21','csv:1')")
        bundle = {"source_id": "official_vic", "state": "VIC", "files_by_role": {
            role: {"spec": {"resource_id": "fixture_" + role}, "file": {"sha256": "0" * 64}}
            for role in ["crash", "vehicle"]}}
        result = _export(db, tmp_path, bundle, definitions, lambda: None)
        assert result[0]["unit_count"] is None
        assert result[0]["canonical_unit_count"] == 1
        assert result[4]["rows"] is None
        unit = json.loads(result[6].read_text())
        assert unit["count_eligible"] is False
        assert unit["quality_reason"] == "definition_unconfirmed"
    finally:
        db.close()
