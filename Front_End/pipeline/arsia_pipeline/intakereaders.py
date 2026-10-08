"""Pure, streaming format readers shared by inspection, sandbox and trusted QA.

Host calls use the trusted parser process boundary; no network, database or model calls live here. Paths are supplied by
the trusted caller or its read-only sandbox manifest; this module never writes
inputs. Source identity and business meaning are deliberately not inferred.
"""
from __future__ import annotations

import csv
from datetime import date, datetime, time
from decimal import Decimal
import json
import math
from pathlib import Path
import re
import zipfile

from .errors import NeedsInput, ValidationFailure, UnsupportedCapability

VERSION = "autonomous-intake-readers-v3"
MAX_COLUMNS = 2048
MAX_SHEETS = 64
MAX_EXPANDED_WORKBOOK = 2 * 1024**3


def _path(file):
    path = Path(file["path"] if isinstance(file, dict) else file)
    if path.is_symlink() or not path.is_file():
        raise ValidationFailure("Input is not a regular non-symlink file.")
    return path


def _header(values):
    if not isinstance(values, (list, tuple)) or not values or len(values) > MAX_COLUMNS:
        raise ValidationFailure("A table needs 1–2048 explicit columns.")
    if any(not isinstance(x, str) or not x.strip() or "\x00" in x for x in values):
        raise ValidationFailure("Column names must be nonempty text.")
    if len(values) != len(set(values)):
        raise ValidationFailure("Duplicate column names need explicit disambiguation.")
    return list(values)


def _kind(file, check_cancelled=lambda: None):
    check_cancelled()
    path = _path(file)
    with path.open("rb") as stream:
        sample = stream.read(65536)
    if sample.startswith(b"%PDF-"):
        return "pdf", sample
    if sample.startswith(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"):
        return "xls", sample
    if sample.startswith(b"PK\x03\x04"):
        from . import parser_guard
        if parser_guard.required():
            return parser_guard.call('format', [file], check_cancelled), sample
        try:
            with zipfile.ZipFile(path) as archive:
                names = archive.namelist()
                return ("xlsx" if "xl/workbook.xml" in names else "zip"), sample
        except zipfile.BadZipFile as exc:
            raise ValidationFailure("Damaged ZIP/workbook archive.") from exc
    stripped = sample.lstrip(b"\xef\xbb\xbf \n\r\t")
    if stripped.startswith((b"{", b"[")):
        return "json", sample
    return "csv", sample


def detect_format(file, check_cancelled=lambda: None):
    return _kind(file, check_cancelled)[0]


def _csv_spec(sample, hints):
    encoding = hints.get("encoding")
    if encoding is None:
        if sample.startswith(b"\xef\xbb\xbf"):
            encoding = "utf-8-sig"
        elif sample.startswith((b"\xff\xfe", b"\xfe\xff")):
            encoding = "utf-16"
        else:
            try:
                # A chunk may end mid-codepoint; only drop that incomplete tail.
                import codecs
                codecs.getincrementaldecoder("utf-8")("strict").decode(sample, final=False)
                encoding = "utf-8"
            except UnicodeError:
                from charset_normalizer import from_bytes
                candidates = [item.encoding for item in from_bytes(sample)][:5]
                raise NeedsInput("CSV encoding is not unambiguously UTF-8; consult its source dictionary before decoding.",
                                 ["Find the publisher's declared export encoding and retry inspection with that encoding."],
                                 {"encoding_candidates": candidates, "requires_encoding_evidence": True})
    if encoding.lower().replace("_", "-") not in {"utf-8", "utf-8-sig", "utf-16", "utf-16-le", "utf-16-be", "cp1252", "windows-1252", "latin-1", "iso-8859-1", "ascii"}:
        raise ValidationFailure("Encoding is outside the supported explicit decoder set.")
    import codecs
    try:
        text = codecs.getincrementaldecoder(encoding)("strict").decode(sample, final=False)
    except UnicodeError as exc:
        raise NeedsInput("File bytes do not match the selected encoding.") from exc
    delimiter = hints.get("delimiter")
    if delimiter is not None and delimiter not in {",", ";", "\t", "|"}:
        raise ValidationFailure("CSV delimiter must be comma, semicolon, tab or pipe.")
    if delimiter is None:
        valid = []
        for candidate in [",", ";", "\t", "|"]:
            try:
                rows = list(csv.reader(text.splitlines(keepends=True), delimiter=candidate, strict=True))[:30]
                # Ignore the potentially incomplete final sampled record.
                stable = rows[:-1] if len(rows) > 1 else rows
                if stable and len(stable[0]) > 1 and all(len(r) == len(stable[0]) for r in stable):
                    valid.append(candidate)
            except csv.Error:
                continue
        if len(valid) > 1:
            raise NeedsInput("Multiple CSV delimiters produce different valid tables.",
                             ["Use the publisher's export specification to resolve the delimiter."], {"delimiter_candidates": valid})
        delimiter = valid[0] if valid else ","
    return {"encoding": encoding, "delimiter": delimiter}


def _excel_value(value, kind=None):
    if kind in {"f", "e"}:
        raise ValidationFailure("Formula/error cells require a values-only source export; formulas are never executed.")
    if value is None or isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValidationFailure("Nonfinite spreadsheet cell.")
        return str(int(value)) if int(value) == value else str(value)
    raise ValidationFailure("Unsupported spreadsheet cell type.")


def _xlsx(path):
    from openpyxl import load_workbook
    with zipfile.ZipFile(path) as archive:
        if len(archive.infolist()) > 10000 or sum(x.file_size for x in archive.infolist()) > MAX_EXPANDED_WORKBOOK:
            raise ValidationFailure("Workbook exceeds the safe expanded-size bound.")
    source = path.open("rb")
    try:
        book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
        book._arsia_source_stream = source
        return book
    except Exception as exc:
        source.close()
        raise ValidationFailure("Damaged, encrypted or unsupported Excel workbook.") from exc


def _close_xlsx(book):
    book.close()
    book._arsia_source_stream.close()


def _json_stream(stream):
    if stream.read(3) != b"\xef\xbb\xbf":
        stream.seek(0)
    return stream


class _UniqueObject(dict):
    def __setitem__(self, key, value):
        if key in self:
            raise ValidationFailure("Duplicate JSON keys cannot be silently overwritten.")
        super().__setitem__(key, value)


def _pairs(pairs):
    result = _UniqueObject()
    for key, value in pairs:
        result[key] = value
    return result


def _json_layout(path, hints=None, check_cancelled=lambda: None):
    """Validate the complete envelope before selecting a streaming row reader.

    Metadata can follow a million features. Keep only bounded metadata and the
    current object-key stack, never the collection. A second streaming pass
    discovers the complete schema; parser hints cannot override native geometry.
    """
    import ijson
    from ijson.common import ObjectBuilder
    hints = hints or {}
    # JSONL is accepted only when each nonempty line is a complete object.
    with path.open("r", encoding="utf-8-sig") as stream:
        first = stream.readline(2 * 1024**2)
        second = stream.readline(2 * 1024**2)
    if second.strip():
        try:
            if isinstance(json.loads(first, object_pairs_hook=_pairs), dict) and isinstance(json.loads(second, object_pairs_hook=_pairs), dict):
                if hints.get('record_path') or hints.get('json_kind', 'jsonl') != 'jsonl':
                    raise ValidationFailure("Parser hints conflict with the JSONL envelope.")
                return {"record_path": None, "json_kind": "jsonl"}
        except ValueError:
            pass
    found, arrays, stack = {}, set(), []
    capture = None
    top_key = None
    root_kind = None
    requested = hints.get('record_path')
    requested_found = False
    metadata_keys = {'type', 'spatialReference', 'crs', 'geometryType'}
    with path.open("rb") as stream:
        try:
            for index, (prefix, event, value) in enumerate(ijson.parse(_json_stream(stream))):
                if index % 1000 == 0:
                    check_cancelled()
                if index == 0:
                    root_kind = event
                if event == 'map_key':
                    keys = stack[-1]
                    if value in keys:
                        raise ValidationFailure("Duplicate JSON keys cannot be silently overwritten.")
                    keys.add(value)
                    if len(keys) > MAX_COLUMNS:
                        raise ValidationFailure("JSON object exceeds the supported field bound.")
                    if len(stack) == 1:
                        top_key = value
                elif len(stack) == 1 and root_kind == 'start_map' and event not in {'end_map', 'end_array'}:
                    if top_key in metadata_keys and capture is None:
                        capture = (top_key, ObjectBuilder(), 0)
                    if event == 'start_array':
                        arrays.add(top_key)
                    elif top_key == 'features':
                        raise ValidationFailure("Feature collection must contain an array of features.")
                if event == 'start_array' and requested == (prefix + '.item' if prefix else 'item'):
                    requested_found = True
                if capture is not None:
                    key, builder, size = capture
                    size += 1 + (len(value) if isinstance(value, str) else 0)
                    if size > 65536:
                        raise ValidationFailure("JSON envelope metadata exceeds its size bound.")
                    builder.event(event, value)
                    capture = (key, builder, size)
                if event in {'start_map', 'start_array'}:
                    stack.append(set() if event == 'start_map' else None)
                    if len(stack) > 64:
                        raise ValidationFailure("JSON nesting exceeds the supported depth.")
                elif event in {'end_map', 'end_array'}:
                    stack.pop()
                if capture is not None and len(stack) == 1:
                    key, builder, _ = capture
                    found[key] = builder.value
                    capture = None
        except (ValueError, UnicodeError, ijson.JSONError) as exc:
            raise ValidationFailure("Invalid JSON input.") from exc
    check_cancelled()
    reference = found.get('spatialReference')
    legacy = found.get('crs')
    if reference is not None and not isinstance(reference, dict):
        raise ValidationFailure("ArcGIS spatialReference must be an object.")
    if 'features' in arrays:
        declared = found.get('type')
        if 'type' in found and declared != 'FeatureCollection':
            raise ValidationFailure("Feature collection has a conflicting type declaration.")
        kind = 'geojson' if declared == 'FeatureCollection' else 'arcgis'
        if kind == 'geojson' and ('spatialReference' in found or 'geometryType' in found):
            raise ValidationFailure("GeoJSON and ArcGIS envelope declarations conflict.")
        if kind == 'arcgis' and 'crs' in found:
            raise ValidationFailure("Legacy GeoJSON CRS cannot define an ArcGIS envelope.")
        if legacy is not None and (not isinstance(legacy, dict) or legacy.get('type') != 'name'
                                  or not isinstance(legacy.get('properties', {}).get('name'), str)):
            raise UnsupportedCapability('LEGACY_CRS_UNSUPPORTED', "Unsupported legacy GeoJSON CRS declaration.")
        spec = {'record_path': 'features.item', 'json_kind': kind,
                'spatial_reference': ((reference or {}).get('latestWkid', (reference or {}).get('wkid', (reference or {}).get('wkt')))
                                      if kind == 'arcgis' else (legacy or {}).get('properties', {}).get('name')),
                'spatial_reference_definition': reference,
                'geojson_default_crs': 'OGC:CRS84' if kind == 'geojson' and 'crs' not in found else None,
                'legacy_crs': legacy, 'geometry_type': found.get('geometryType')}
    elif found.get('type') == 'FeatureCollection':
        raise ValidationFailure("GeoJSON FeatureCollection lacks its features array.")
    elif root_kind == 'start_array':
        spec = {'record_path': 'item', 'json_kind': 'records'}
    elif requested and requested_found:
        spec = {'record_path': requested, 'json_kind': 'records'}
    else:
        candidates = arrays & {'records', 'data'}
        if len(candidates) > 1:
            raise NeedsInput("Multiple JSON record arrays require explicit selection.")
        if not candidates:
            raise NeedsInput("This JSON object is a document, not a recognized record collection.")
        spec = {'record_path': next(iter(candidates)) + '.item', 'json_kind': 'records'}
    if requested and requested != spec['record_path'] or hints.get('json_kind', spec['json_kind']) != spec['json_kind']:
        raise ValidationFailure("Parser hints conflict with the observed JSON envelope.")
    return spec


def _clean_json(value):
    if isinstance(value, Decimal):
        return str(value)
    if isinstance(value, dict):
        return {k: _clean_json(v) for k, v in value.items()}
    if isinstance(value, list):
        return [_clean_json(v) for v in value]
    return value


def _json_rows(path, spec):
    import ijson
    if spec.get("json_kind") == "jsonl":
        with path.open("r", encoding="utf-8-sig") as stream:
            for number, line in enumerate(stream, 1):
                if not line.strip():
                    raise ValidationFailure("Empty JSONL records cannot silently disappear.")
                try:
                    def invalid_constant(value):
                        raise ValidationFailure("Nonfinite JSON constants are invalid.")
                    yield number, json.loads(line, parse_float=Decimal, object_pairs_hook=_pairs, parse_constant=invalid_constant)
                except ValueError as exc:
                    raise ValidationFailure("Invalid JSONL record.") from exc
    else:
        with path.open("rb") as stream:
            try:
                for number, item in enumerate(ijson.items(_json_stream(stream), spec["record_path"], map_type=_UniqueObject), 1):
                    yield number, item
            except (ValueError, ijson.JSONError) as exc:
                raise ValidationFailure("Invalid JSON record stream.") from exc


def _feature_row(item, spec):
    """Preserve geometry; only a validated native point exposes x/y columns."""
    geojson = spec['json_kind'] == 'geojson'
    if geojson and 'crs' in item:
        raise UnsupportedCapability('FEATURE_CRS_UNSUPPORTED', 'Feature-level legacy CRS requires an explicit supported geometry plan.')
    if geojson and item.get('type') != 'Feature':
        raise ValidationFailure("GeoJSON collection member is not a Feature.")
    row = item.get('properties' if geojson else 'attributes')
    if geojson and 'properties' in item and row is None:
        row = {}
    if not isinstance(row, dict) or any(name in row for name in ('__geometry', '__geometry_x', '__geometry_y')):
        raise ValidationFailure("Invalid feature attributes or reserved field collision.")
    if geojson and 'geometry' not in item:
        raise ValidationFailure("GeoJSON Feature lacks its geometry member.")
    geometry = item.get('geometry')
    x = y = None
    if geometry is not None:
        if not isinstance(geometry, dict):
            raise ValidationFailure("Feature geometry must be an object or null.")
        if geojson:
            if 'crs' in geometry or 'spatialReference' in geometry:
                raise UnsupportedCapability('FEATURE_CRS_UNSUPPORTED', 'Per-geometry legacy CRS cannot be interpreted as RFC 7946.')
            kind = geometry.get('type')
            if kind not in {'Point', 'MultiPoint', 'LineString', 'MultiLineString', 'Polygon', 'MultiPolygon', 'GeometryCollection'}:
                raise ValidationFailure("Unknown GeoJSON geometry type.")
            point = kind == 'Point'
            coords = geometry.get('coordinates')
        else:
            point = 'x' in geometry or 'y' in geometry
            coords = [geometry.get('x'), geometry.get('y')]
            if point and any(key in geometry for key in ('rings', 'paths', 'points')):
                raise ValidationFailure("ArcGIS point conflicts with non-point geometry.")
            if point and spec.get('geometry_type') not in (None, 'esriGeometryPoint'):
                raise ValidationFailure("ArcGIS geometry conflicts with its declared type.")
        if point:
            if not isinstance(coords, list) or len(coords) not in (2, 3):
                raise ValidationFailure("Point geometry needs two or three numeric ordinates.")
            if any(isinstance(v, bool) or not isinstance(v, (int, float, Decimal)) or not math.isfinite(v) for v in coords):
                raise ValidationFailure("Point ordinates must be finite JSON numbers.")
            x, y = coords[:2]
            if geojson and spec.get('geojson_default_crs') and not (-180 <= x <= 180 and -90 <= y <= 90):
                raise ValidationFailure("RFC 7946 point is outside longitude/latitude bounds.")
    return {**row, '__geometry': geometry, '__geometry_x': x, '__geometry_y': y}


def detect_tables(file, hints=None, check_cancelled=lambda: None):
    """Describe tables by actual format; hints resolve documented ambiguity."""
    from .parser_guard import required, call
    if required(): return call("detect", [file, hints], check_cancelled)
    path, hints = _path(file), hints or {}
    kind, sample = _kind(file)
    if kind == "csv":
        spec = _csv_spec(sample, hints)
        try:
            with path.open("r", encoding=spec["encoding"], newline="") as stream:
                header = _header(next(csv.reader(stream, delimiter=spec["delimiter"], strict=True), None))
        except (UnicodeError, csv.Error) as exc:
            raise ValidationFailure("CSV header cannot be decoded under its declared parser.") from exc
        return [{"table_id": "default", "format": "csv", "header": header, **spec}]
    if kind == "xlsx":
        from .workbook_plan import describe_workbook, input_digest, validate_hint
        file_sha256 = input_digest(path, file, check_cancelled)
        book = _xlsx(path)
        try:
            if len(book.worksheets) > MAX_SHEETS:
                raise ValidationFailure("Workbook has too many worksheets.")
            result = describe_workbook(book, file_sha256, _header, _excel_value, check_cancelled)
            selected = hints.get('sheet', hints.get('table_id'))
            if selected is not None and selected not in {t['table_id'] for t in result}:
                raise ValidationFailure('Workbook sheet is absent from the actual input.')
            if selected is None and len(result) != 1 and any(k in hints for k in ('header_row', 'header', 'parser_plan')):
                raise ValidationFailure('A workbook parser plan must identify its specific sheet.')
            for table in result:
                if table['table_id'] == selected or selected is None:
                    validate_hint(table, hints)
            return result
        finally:
            _close_xlsx(book)
    if kind == "xls":
        if hints.get('header_row', 1) != 1 or 'parser_plan' in hints:
            raise UnsupportedCapability('XLS_HEADER_PLAN_UNSUPPORTED', 'The bounded workbook header plan currently supports XLSX; legacy XLS requires its existing first-row header.')
        import xlrd
        try:
            book = xlrd.open_workbook(path, on_demand=True)
        except Exception as exc:
            raise ValidationFailure("Damaged, encrypted or unsupported legacy XLS workbook.") from exc
        try:
            if book.nsheets > MAX_SHEETS:
                raise ValidationFailure("Workbook has too many worksheets.")
            return [{"table_id": sheet.name, "sheet": sheet.name, "format": "xls", "header": _header(sheet.row_values(0)),
                     "date_mode": book.datemode} for sheet in book.sheets() if sheet.nrows]
        finally:
            book.release_resources()
    if kind == "json":
        spec = _json_layout(path, hints, check_cancelled)
        header = []
        number = 0
        geometry_types, geometry_references = set(), {}
        for number, record in _json_rows(path, spec):
            if number % 1000 == 0:
                check_cancelled()
            if not isinstance(record, dict):
                raise UnsupportedCapability('ARRAY_RECORDS_UNSUPPORTED', "JSON array records require a reviewed positional table plan.")
            if spec["json_kind"] in {"arcgis", "geojson"}:
                native = _feature_row(record, spec)
                geometry = native['__geometry']
                geometry_types.add('null' if geometry is None else geometry.get('type',
                                   'Point' if 'x' in geometry and 'y' in geometry else
                                   'Polygon' if 'rings' in geometry else 'LineString' if 'paths' in geometry else
                                   'MultiPoint' if 'points' in geometry else 'Unknown'))
                if geometry is not None and 'spatialReference' in geometry:
                    reference = _clean_json(geometry['spatialReference'])
                    key = json.dumps(reference, sort_keys=True)
                    geometry_references.setdefault(key, {'reference': reference, 'first_row': number, 'row_count': 0})['row_count'] += 1
                    if len(geometry_references) > 32:
                        raise UnsupportedCapability('MULTIPLE_CRS_PLAN_UNSUPPORTED', 'Too many per-geometry spatial references for a single geometry plan.')
                record = {k: v for k, v in native.items()
                          if k not in {'__geometry', '__geometry_x', '__geometry_y'}}
            for name in record:
                if name not in header:
                    header.append(name)
                    if len(header) > MAX_COLUMNS:
                        raise ValidationFailure("JSON table exceeds the supported column bound.")
        if spec["json_kind"] in {"arcgis", "geojson"}:
            if any(name in header for name in ("__geometry", "__geometry_x", "__geometry_y")):
                raise ValidationFailure("Native field collides with reserved geometry envelope.")
            header.extend(["__geometry", "__geometry_x", "__geometry_y"])
            spec.update(observed_geometry_types=sorted(geometry_types), geometry_spatial_references=list(geometry_references.values()))
        return [{"table_id": "default", "format": "json", "header": _header(header),
                 "header_complete": True, "header_inspected_records": number, **spec}]
    if kind in {"pdf", "zip"}:
        return []
    raise UnsupportedCapability('FILE_FORMAT_UNSUPPORTED', "Unsupported file format.")


def iter_table(file, spec=None, check_cancelled=lambda: None):
    """Yield full native rows; never infer business grain, identifiers or dates."""
    from .parser_guard import required, packets
    if required():
        for locator, row in packets("rows", [file, spec], check_cancelled): yield locator, row
        return
    path = _path(file)
    if detect_format(file) == 'xlsx':
        # Recompute even if a caller supplies a header or a complete plan. The
        # immutable input determines which rows can be metadata, not the adapter.
        observed = detect_tables(file, spec, check_cancelled)
        selected = (spec or {}).get('sheet', (spec or {}).get('table_id'))
        matches = [t for t in observed if t['table_id'] == selected or (selected is None and len(observed) == 1)]
        if len(matches) != 1:
            raise NeedsInput('Choose a documented worksheet/table before reading records.')
        spec = matches[0]
        if spec['empty']:
            raise ValidationFailure('An empty worksheet cannot be read as a fact table.')
    elif spec is None:
        tables = detect_tables(file, check_cancelled=check_cancelled)
        if len(tables) != 1:
            raise NeedsInput("Choose a documented worksheet/table before reading records.")
        spec = tables[0]
    elif "header" not in spec:
        observed = detect_tables(file, spec, check_cancelled)
        matches = [t for t in observed if spec.get("sheet", spec.get("table_id", t["table_id"])) == t["table_id"]]
        if len(matches) != 1:
            raise NeedsInput("The source contract does not identify exactly one native table.")
        spec = {**matches[0], **spec}
    kind = spec["format"]
    expected = spec["header"]
    check_cancelled()
    if kind == "csv":
        try:
            with path.open("r", encoding=spec.get("encoding", "utf-8-sig"), newline="") as stream:
                reader = csv.reader(stream, delimiter=spec.get("delimiter", ","), strict=True)
                if _header(next(reader, None)) != expected:
                    raise ValidationFailure("CSV header differs from the inspected parser contract.")
                for number, row in enumerate(reader, 1):
                    if number % 1000 == 0:
                        check_cancelled()
                    if len(row) != len(expected):
                        raise ValidationFailure("CSV record width differs from its header.", details={"row_locator": f"csv:{number}"})
                    yield f"csv:{number}", dict(zip(expected, row, strict=True))
        except (UnicodeError, csv.Error) as exc:
            raise ValidationFailure("CSV decoding/quoting failed under the selected parser.") from exc
    elif kind == "xlsx":
        book = _xlsx(path)
        try:
            sheet = book[spec["sheet"]]
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            for _ in range(spec['header_row'] - 1):
                check_cancelled()
                next(rows)
            cells = list(next(rows, ()))
            while cells and cells[-1].value is None:
                cells.pop()
            if _header([_excel_value(c.value, c.data_type) for c in cells]) != expected:
                raise ValidationFailure("Workbook header changed after inspection.")
            blanks = []
            for number, cells in enumerate(rows, spec['header_row'] + 1):
                if number % 1000 == 0:
                    check_cancelled()
                if not cells or all(c.value is None for c in cells):
                    # Store only the interval endpoints, not a huge blank tail.
                    blanks = [blanks[0] if blanks else number, number]
                    continue
                if blanks:
                    for blank in range(blanks[0], blanks[1] + 1):
                        if blank % 1000 == 0:
                            check_cancelled()
                        yield json.dumps([sheet.title, blank], separators=(",", ":")), dict.fromkeys(expected)
                    blanks = []
                if len(cells) > len(expected) and any(c.value is not None for c in cells[len(expected):]):
                    raise ValidationFailure("Workbook row has undeclared extra columns.")
                values = [_excel_value(c.value, c.data_type) for c in cells[:len(expected)]]
                values.extend([None] * (len(expected) - len(values)))
                yield json.dumps([sheet.title, number], separators=(",", ":")), dict(zip(expected, values, strict=True))
        finally:
            _close_xlsx(book)
    elif kind == "xls":
        import xlrd
        book = xlrd.open_workbook(path, on_demand=True)
        try:
            sheet = book.sheet_by_name(spec["sheet"])
            if _header(sheet.row_values(0)) != expected:
                raise ValidationFailure("XLS header changed after inspection.")
            for number in range(1, sheet.nrows):
                if number % 1000 == 0:
                    check_cancelled()
                values = []
                for cell in sheet.row(number):
                    if cell.ctype == xlrd.XL_CELL_ERROR:
                        raise ValidationFailure("XLS error cell requires an original values-only export.")
                    if cell.ctype == xlrd.XL_CELL_DATE:
                        value = xlrd.xldate_as_datetime(cell.value, book.datemode).isoformat()
                    elif cell.ctype in {xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK}:
                        value = None
                    else:
                        value = _excel_value(cell.value)
                    values.append(value)
                yield json.dumps([sheet.name, number + 1], separators=(",", ":")), dict(zip(expected, values, strict=True))
        finally:
            book.release_resources()
    elif kind == "json":
        # Caller-provided headers do not authorize another parser or geometry
        # interpretation. Recheck the envelope even for previously inspected specs.
        observed = _json_layout(path, spec, check_cancelled)
        spec = {**spec, **observed}
        for number, item in _json_rows(path, spec):
            if number % 1000 == 0:
                check_cancelled()
            if not isinstance(item, dict):
                raise ValidationFailure("JSON record is not an object.")
            if spec.get("json_kind") in {"arcgis", "geojson"}:
                row = _feature_row(item, spec)
            else:
                row = item
            yield f"json:{number}", _clean_json(row)
    else:
        raise UnsupportedCapability('TABLE_READER_UNSUPPORTED', "Selected resource is not a supported tabular file.")
    check_cancelled()
