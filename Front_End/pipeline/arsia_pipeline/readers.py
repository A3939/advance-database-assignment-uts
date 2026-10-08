"""Bounded native readers. Input bytes are opened read-only and never repaired.

This is a new local-test adapter, informed by Role B / Peixian's original
native-intake contract. It does not impersonate the course parser or its FP1.
"""
from __future__ import annotations

import csv
import json
from datetime import date, datetime, time
from decimal import Decimal
from pathlib import Path
import math
import zipfile

from .errors import NeedsInput, ValidationFailure

READER_VERSION = "local-native-reader-v1"
MAX_COLUMNS = 1024
MAX_SHEETS = 32
MAX_UNCOMPRESSED_XLSX = 4 * 1024**3


def _path(file):
    value = Path(file["path"] if isinstance(file, dict) else file)
    if value.is_symlink() or not value.is_file():
        raise ValidationFailure("Input must be an immutable regular file.")
    return value


def _kind(path):
    with path.open("rb") as stream:
        sample = stream.read(8192)
    if sample.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(path) as archive:
                if "xl/workbook.xml" not in archive.namelist():
                    raise NeedsInput("The ZIP attachment is not an Excel workbook.")
                if len(archive.infolist()) > 10000 or sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED_XLSX:
                    raise ValidationFailure("Workbook exceeds the safe expanded-size limit.")
        except zipfile.BadZipFile as exc:
            raise ValidationFailure("Invalid workbook archive.") from exc
        return "xlsx"
    if sample.startswith(b"\xd0\xcf\x11\xe0"):
        raise NeedsInput("Legacy .xls needs an explicitly reviewed .xlsx or CSV export.")
    if b"\x00" in sample:
        raise NeedsInput("Binary or UTF-16 input is not a supported UTF-8 CSV.")
    return "csv"


def _header(values):
    if not values or len(values) > MAX_COLUMNS:
        raise ValidationFailure("The table has no header or too many columns.")
    if any(not isinstance(v, str) or not v.strip() or "\x00" in v for v in values):
        raise ValidationFailure("Headers must be nonempty text; no columns are invented.")
    if len(values) != len(set(values)):
        raise ValidationFailure("Duplicate column names make field mapping ambiguous.")
    return list(values)


def _workbook(path):
    from openpyxl import load_workbook

    try:
        return load_workbook(path, read_only=True, data_only=False, keep_links=False)
    except Exception as exc:
        raise ValidationFailure("The Excel workbook could not be parsed safely.") from exc


def inspect_file(file):
    from .parser_guard import required, call
    if required(): return call("native_detect", [file])
    """Return table descriptions without reading personal data into a profile.

    Column and sheet identities are exact. CSV encoding is strict UTF-8, with
    optional BOM; no automatic lossy decoding or delimiter guessing is used.
    """
    path = _path(file)
    kind = _kind(path)
    if kind == "csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                header = next(csv.reader(stream, strict=True), None)
        except (UnicodeError, csv.Error) as exc:
            raise NeedsInput("CSV requires valid UTF-8 and comma-separated fields.") from exc
        return [{"format": "csv", "encoding": "utf-8", "sheet": None, "header": _header(header)}]
    book = _workbook(path)
    try:
        if len(book.worksheets) > MAX_SHEETS:
            raise ValidationFailure("Workbook contains too many sheets.")
        tables = []
        for sheet in book.worksheets:
            sheet.reset_dimensions()
            first = next(sheet.iter_rows(), ())
            if not first or all(c.value is None for c in first):
                continue
            if any(c.data_type == "f" for c in first):
                raise ValidationFailure("Formula headers are not accepted.")
            tables.append({"format": "xlsx", "encoding": None, "sheet": sheet.title,
                           "header": _header([c.value for c in first])})
        if not tables:
            raise NeedsInput("Workbook has no nonempty first-row table headers.")
        return tables
    finally:
        book.close()


def _text(cell, locator):
    value = cell.value
    if cell.data_type in {"f", "e"}:
        raise ValidationFailure("Excel formulas and error cells require a reviewed values-only export.",
                                details={"row_locator": locator, "cell": cell.coordinate})
    if value is None or isinstance(value, str):
        if isinstance(value, str) and "\x00" in value:
            raise ValidationFailure("NUL characters are not accepted in native values.")
        return value
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, (int, float, Decimal)):
        if isinstance(value, float) and not math.isfinite(value):
            raise ValidationFailure("Nonfinite numeric cell.")
        # Integral numeric cells have no hidden leading-zero string identity.
        return str(int(value)) if value == int(value) else str(value)
    raise ValidationFailure("Unsupported Excel cell value type.", details={"row_locator": locator})


def iter_rows(file, table=None, check_cancelled=lambda: None):
    from .parser_guard import required, packets
    if required():
        for locator, row in packets("native_rows", [file, table], check_cancelled): yield locator, row
        return
    """Yield (physical/logical row locator, exact native text-or-null payload)."""
    path = _path(file)
    if table is None:
        tables = inspect_file(file)
        if len(tables) != 1:
            raise NeedsInput("Select exactly one worksheet in the reviewed profile.")
        table = tables[0]
    expected = table["header"]
    if table["format"] == "csv":
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as stream:
                reader = csv.reader(stream, strict=True)
                if _header(next(reader, None)) != expected:
                    raise ValidationFailure("Header changed between inspection and reading.")
                # B's csv-logical-v1 counts data records from one, excluding
                # the header (and irrespective of quoted embedded newlines).
                for number, cells in enumerate(reader, 1):
                    if number % 1000 == 0:
                        check_cancelled()
                    if len(cells) != len(expected):
                        raise ValidationFailure("CSV record width differs from the header.",
                                                details={"row_locator": f"csv:{number}"})
                    if any("\x00" in value for value in cells):
                        raise ValidationFailure("CSV contains a NUL character.")
                    yield f"csv:{number}", dict(zip(expected, cells, strict=True))
        except (UnicodeError, csv.Error) as exc:
            raise ValidationFailure("CSV contains invalid encoding or quoting.") from exc
    elif table["format"] == "xlsx":
        book = _workbook(path)
        try:
            if table["sheet"] not in book.sheetnames:
                raise ValidationFailure("Selected worksheet is missing.")
            sheet = book[table["sheet"]]
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            if _header([c.value for c in next(rows, ())]) != expected:
                raise ValidationFailure("Worksheet header changed between inspection and reading.")
            pending_blank_start = None
            pending_blank_count = 0
            for number, cells in enumerate(rows, 2):
                if number % 1000 == 0:
                    check_cancelled()
                if all(c.value is None for c in cells):
                    if pending_blank_start is None:
                        pending_blank_start = number
                    pending_blank_count += 1
                    continue
                # Interior blank rows remain records and fail required-key QA;
                # only the final empty tail of a worksheet is discarded.
                if pending_blank_count:
                    for empty_row in range(pending_blank_start, pending_blank_start + pending_blank_count):
                        if empty_row % 1000 == 0:
                            check_cancelled()
                        yield json.dumps([table["sheet"], empty_row], separators=(",", ":")), dict.fromkeys(expected)
                    pending_blank_start, pending_blank_count = None, 0
                locator = json.dumps([table["sheet"], number], separators=(",", ":"))
                if len(cells) > len(expected) and any(c.value is not None for c in cells[len(expected):]):
                    raise ValidationFailure("Worksheet row width differs from its header.", details={"row_locator": locator})
                values = [_text(c, locator) for c in cells[:len(expected)]]
                values.extend([None] * (len(expected) - len(values)))
                yield locator, dict(zip(expected, values, strict=True))
        finally:
            book.close()
    else:
        raise NeedsInput("Unsupported table format.")
    check_cancelled()
